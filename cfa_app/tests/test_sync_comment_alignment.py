"""Scenario tests for row-sync + comment alignment + write-by-key.

Run:  .venv/Scripts/python.exe cfa_app/tests/test_sync_comment_alignment.py

Uses the local master workbook as a fixture (client data, gitignored). If it isn't present the
suite skips with a clear message rather than failing.
"""

from __future__ import annotations

import io
import pathlib
import sys

ROOT = pathlib.Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "cfa_app"))

import openpyxl                                              # noqa: E402
from openpyxl.comments import Comment                       # noqa: E402

from app.core.detection import norm_key                     # noqa: E402
from app.core.engine import (                               # noqa: E402
    _ENTITY_CELL_RE, _ws_getter, clear_separator_entity_data, find_entity_header_row,
    find_entity_row_and_col, find_target_header, find_template_period_sheet, load_write_workbook,
    save_workbook_to_bytes, sync_template_rows, transfer_into_master,
)
from app.core.models import DetectionConfig, SourceData     # noqa: E402

MASTER = ROOT / "CFA verification_P5 2026.xlsx"
cfg = DetectionConfig()
_results: list[tuple[str, bool, str]] = []


def check(name: str, ok: bool, detail: str = "") -> None:
    _results.append((name, ok, detail))
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail and not ok else ""))


def _keyseq(ws):
    g = _ws_getter(ws); mc = min(ws.max_column or 1, 400)
    hr, fc, lc = find_target_header(g, cfg.header_scan_rows, mc, cfg)
    out = []
    for r in range(hr + 1, (ws.max_row or hr) + 1):
        f, l = norm_key(g(r, fc)), norm_key(g(r, lc))
        if f is None and l is None:
            continue
        out.append((f, l))
    return out


def _key_of(ws, r):
    return (norm_key(ws.cell(row=r, column=1).value), norm_key(ws.cell(row=r, column=2).value))


def _comment_keys(ws):
    """Map comment text -> the (FORM, LINE) of the row it sits on."""
    out = {}
    for row in ws.iter_rows():
        for c in row:
            if c.comment and c.comment.text:
                out[c.comment.text.split(":")[0].strip()] = _key_of(ws, c.row)
    return out


def scenario_missing_rows_and_comment_alignment(mb):
    """Year file missing template rows + comments below the gap -> sync adds rows, comments stay put."""
    print("Scenario 1: sync adds missing rows, comments stay on their (FORM, LINE)")
    values_wb = openpyxl.load_workbook(io.BytesIO(mb), data_only=True)
    tname = find_template_period_sheet(values_wb)
    tpl_seq = _keyseq(values_wb[tname])

    yb = load_write_workbook(mb)
    for s in list(yb.sheetnames):
        if s not in (tname, "Methodology"):
            del yb[s]
    ws = yb[tname]; g = _ws_getter(ws); mc = min(ws.max_column or 1, 400)
    hr, fc, lc = find_target_header(g, cfg.header_scan_rows, mc, cfg)
    keyed = [r for r in range(hr + 1, (ws.max_row or hr) + 1)
             if not (norm_key(g(r, fc)) is None and norm_key(g(r, lc)) is None)]
    tag_rows = [keyed[40], keyed[70], keyed[100]]
    want = {}
    for i, r in enumerate(tag_rows):
        ws.cell(row=r, column=4).comment = Comment(f"TAG{i}", "t")
        want[f"TAG{i}"] = _key_of(ws, r)
    for dr in sorted([keyed[5], keyed[25]], reverse=True):
        ws.delete_rows(dr, 1)
    year_bytes = save_workbook_to_bytes(yb)

    wb = load_write_workbook(year_bytes)
    added = sync_template_rows(values_wb, wb, cfg, only_periods={int(tname[1:].split()[0])})
    out = save_workbook_to_bytes(wb, source_bytes=year_bytes, changed_sheets=set(added))
    w = openpyxl.load_workbook(io.BytesIO(out))

    check("rows were synced", bool(added), f"added={added}")
    check("keyed sequence matches template", _keyseq(w[tname]) == tpl_seq)
    got = _comment_keys(w[tname])
    for t, k in want.items():
        check(f"comment {t} stays on its (FORM,LINE)", got.get(t) == k, f"want {k} got {got.get(t)}")


def scenario_no_change_preserves_comments(mb):
    """Aligned sheet, no inserts -> comments preserved exactly (transplant path)."""
    print("Scenario 2: aligned sheet (no structural change) preserves comments")
    values_wb = openpyxl.load_workbook(io.BytesIO(mb), data_only=True)
    tname = find_template_period_sheet(values_wb)
    yb = load_write_workbook(mb)
    for s in list(yb.sheetnames):
        if s not in (tname, "Methodology"):
            del yb[s]
    before = sum(1 for row in yb[tname].iter_rows() for c in row if c.comment)
    year_bytes = save_workbook_to_bytes(yb)
    wb = load_write_workbook(year_bytes)
    added = sync_template_rows(values_wb, wb, cfg, only_periods={int(tname[1:].split()[0])})
    out = save_workbook_to_bytes(wb, source_bytes=year_bytes, changed_sheets=set(added))
    after = sum(1 for row in openpyxl.load_workbook(io.BytesIO(out))[tname].iter_rows()
                for c in row if c.comment)
    check("no rows added on an aligned sheet", not added, f"added={added}")
    check("comment count preserved", before == after, f"before={before} after={after}")


def scenario_write_by_key_on_shifted_sheet(mb):
    """A sheet whose header is shifted down -> values land on the sheet's own (FORM,LINE) rows."""
    print("Scenario 3: shifted sheet -> values land on correct rows (write-by-key)")
    values_wb = openpyxl.load_workbook(io.BytesIO(mb), data_only=True)
    tname = find_template_period_sheet(values_wb)
    period = int(tname[1:].split()[0])
    wb = load_write_workbook(mb)
    for s in list(wb.sheetnames):
        if s not in (tname, "Methodology"):
            del wb[s]
    wb[tname].insert_rows(1, 1)                     # shift the whole sheet down one row
    ws = wb[tname]; g = _ws_getter(ws); mc = min(ws.max_column or 1, 400)
    er = find_entity_header_row(g, cfg.entity_row_scan, mc)
    hr, fc, lc = find_target_header(g, cfg.header_scan_rows, mc, cfg)
    keyed = [r for r in range(hr + 1, (ws.max_row or hr) + 1)
             if not (norm_key(g(r, fc)) is None and norm_key(g(r, lc)) is None)]
    lines = [(g(keyed[k], fc), g(keyed[k], lc)) for k in (10, 30, 60)]
    src = SourceData(filename="x", entity="001", month=period, year=2026, sheet_name=tname,
                     header_row=hr, form_col=fc, line_col=lc, check_col=3,
                     rows=[(a, b, "NOT EQUAL", 0.0) for a, b in lines],
                     currency="USD", company="", source_url="https://sp/x.xlsx")
    fr = transfer_into_master(values_wb, wb, src, cfg, region=None)
    go = _ws_getter(ws); _, ecol = find_entity_row_and_col(go, cfg.entity_row_scan, mc, "001")
    check("entity row detected on the shifted sheet", er == 9, f"entity_row={er}")
    ok = True
    for a, b in lines:
        row = next(r for r in range(hr + 1, (ws.max_row or hr) + 1)
                   if (norm_key(g(r, fc)), norm_key(g(r, lc))) == (norm_key(a), norm_key(b)))
        val = ws.cell(row=row, column=ecol).value
        label_ok = ws.cell(row=row, column=1).value not in (None, "")
        ok = ok and val == "NOT EQUAL" and label_ok
    check("all values on labelled (FORM,LINE) rows", ok, f"written={fr.written}")
    # hyperlink on the entity row, not the row above
    link = ws.cell(row=er, column=ecol).hyperlink
    above = ws.cell(row=er - 1, column=ecol).hyperlink
    check("hyperlink on entity row, not the row above", bool(link) and not above)


def scenario_file_validity(mb):
    """After a synced run the workbook re-opens and every XML part parses."""
    print("Scenario 4: output file stays valid (reload + XML parse)")
    import xml.etree.ElementTree as ET
    import zipfile
    values_wb = openpyxl.load_workbook(io.BytesIO(mb), data_only=True)
    tname = find_template_period_sheet(values_wb)
    yb = load_write_workbook(mb)
    for s in list(yb.sheetnames):
        if s not in (tname, "Methodology"):
            del yb[s]
    ws = yb[tname]; g = _ws_getter(ws); mc = min(ws.max_column or 1, 400)
    hr, fc, lc = find_target_header(g, cfg.header_scan_rows, mc, cfg)
    keyed = [r for r in range(hr + 1, (ws.max_row or hr) + 1)
             if not (norm_key(g(r, fc)) is None and norm_key(g(r, lc)) is None)]
    ws.delete_rows(keyed[10], 1)
    year_bytes = save_workbook_to_bytes(yb)
    wb = load_write_workbook(year_bytes)
    added = sync_template_rows(values_wb, wb, cfg, only_periods={int(tname[1:].split()[0])})
    out = save_workbook_to_bytes(wb, source_bytes=year_bytes, changed_sheets=set(added))
    try:
        openpyxl.load_workbook(io.BytesIO(out)); reload_ok = True
    except Exception:
        reload_ok = False
    z = zipfile.ZipFile(io.BytesIO(out))
    bad = []
    for p in z.namelist():
        if p.endswith((".xml", ".rels")):
            try:
                ET.fromstring(z.read(p))
            except Exception as e:
                bad.append((p, str(e)))
    check("output re-opens in openpyxl", reload_ok)
    check("every XML part parses", not bad, f"bad={bad[:3]}")


def scenario_clear_separator_rows(mb):
    """Stale check values on a blank-FORM/LINE row are removed; labelled rows are untouched."""
    print("Scenario 5: clear stale values from separator rows, keep labelled rows")
    values_wb = openpyxl.load_workbook(io.BytesIO(mb), data_only=True)
    tname = find_template_period_sheet(values_wb)
    period = int(tname[1:].split()[0])
    wb = load_write_workbook(mb)
    for s in list(wb.sheetnames):
        if s not in (tname, "Methodology"):
            del wb[s]
    ws = wb[tname]; g = _ws_getter(ws); mc = min(ws.max_column or 1, 400)
    hr, fc, lc = find_target_header(g, cfg.header_scan_rows, mc, cfg)
    er = find_entity_header_row(g, cfg.entity_row_scan, mc)
    ecol = next(c for c in range(1, mc + 1) if _ENTITY_CELL_RE.match(str(g(er, c) or "").strip()))
    sep_row = next(r for r in range(hr + 1, (ws.max_row or hr) + 1)
                   if norm_key(g(r, fc)) is None and norm_key(g(r, lc)) is None)
    lab_row = next(r for r in range(hr + 1, (ws.max_row or hr) + 1)
                   if norm_key(g(r, fc)) is not None or norm_key(g(r, lc)) is not None)
    ws.cell(row=sep_row, column=ecol).value = "OK"        # stale value on a separator row
    ws.cell(row=lab_row, column=ecol).value = "NOT EQUAL"  # legit value on a labelled row
    cleared = clear_separator_entity_data(wb, cfg, only_periods={period})
    check("separator-row value was cleared", ws.cell(row=sep_row, column=ecol).value is None,
          f"cleared={cleared}")
    check("labelled-row value kept", ws.cell(row=lab_row, column=ecol).value == "NOT EQUAL")


def main():
    if not MASTER.exists():
        print(f"SKIP: fixture not found: {MASTER}")
        return 0
    mb = MASTER.read_bytes()
    scenario_missing_rows_and_comment_alignment(mb)
    scenario_no_change_preserves_comments(mb)
    scenario_write_by_key_on_shifted_sheet(mb)
    scenario_clear_separator_rows(mb)
    scenario_file_validity(mb)
    passed = sum(1 for _, ok, _ in _results if ok)
    total = len(_results)
    print(f"\n==== {passed}/{total} checks passed ====")
    return 0 if passed == total else 1


if __name__ == "__main__":
    raise SystemExit(main())
